"""Reads the attendance grid (read-only) and classifies each cell.

Simple rules on purpose: the sheet format is fixed by the Corps, and the cell
text is free-form, so we only decide "responded or not" plus a coarse status.
"""
from __future__ import annotations
from dataclasses import dataclass
import re

MISSING, ATTENDING, NOT_ATTENDING, LATE, SEE_OTHER, OTHER = (
    "MISSING", "ATTENDING", "NOT_ATTENDING", "LATE", "SEE_OTHER", "OTHER")


@dataclass
class Response:
    status: str
    raw: str
    reason: str = ""

    @property
    def responded(self) -> bool:
        return self.status != MISSING


def classify(cell) -> Response:
    raw = "" if cell is None else str(cell).strip()
    low = raw.lower()
    if not raw:
        return Response(MISSING, raw)
    if re.match(r"^(yes|y)\b", low):
        return Response(ATTENDING, raw)
    if re.match(r"^late\b", low):
        return Response(LATE, raw, raw[4:].strip(" ,:-"))
    if re.match(r"^(no|n)\b", low):
        return Response(NOT_ATTENDING, raw, re.sub(r"^(no|n)\b[\s,:-]*", "", raw, flags=re.I))
    if low.startswith("like "):
        # "like ethan" = same answer as another cadet; counted as a response
        return Response(SEE_OTHER, raw, raw[5:].strip())
    return Response(OTHER, raw)


@dataclass
class CadetRow:
    name: str              # "Last, First" — the key used in contacts.yaml
    rank: str
    role: str              # PL / PSG / SL / ... from the first column, if any
    group: str             # section header above the row, e.g. "Freshman Platoon"
    responses: dict        # column header -> Response


def parse_grid(rows: list[list], columns: dict, event_columns: list[str]) -> list[CadetRow]:
    """rows: the worksheet as a list of lists of cell values."""
    want = {k: v.strip().lower() for k, v in columns.items()}
    header_idx, colmap = None, {}
    for i, row in enumerate(rows):
        cells = [("" if c is None else str(c)).strip().lower() for c in row]
        if want["last"] in cells and want["first"] in cells:
            header_idx = i
            colmap = {c: j for j, c in enumerate(cells) if c}
            break
    if header_idx is None:
        raise ValueError(f"Header row with {columns['last']!r}/{columns['first']!r} not found")

    missing_cols = [e for e in event_columns if e.strip().lower() not in colmap]
    if missing_cols:
        raise ValueError(f"Event columns not found in header: {missing_cols}")

    def cell(row, key):
        j = colmap.get(key)
        return "" if j is None or j >= len(row) or row[j] is None else str(row[j]).strip()

    out, group = [], ""
    for row in rows[header_idx + 1:]:
        first_col = "" if not row or row[0] is None else str(row[0]).strip()
        last, first = cell(row, want["last"]), cell(row, want["first"])
        if not last and not first:
            if first_col:
                group = first_col          # section header row
            continue
        out.append(CadetRow(
            name=f"{last}, {first}",
            rank=cell(row, want["rank"]),
            role=first_col,
            group=group,
            responses={e: classify(cell(row, e.strip().lower())) for e in event_columns},
        ))
    return out


# ---------- sources ----------

def read_rows(source: dict) -> list[list]:
    kind = source.get("type", "xlsx")
    if kind == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(source["path"], read_only=True, data_only=True)
        ws = wb[source["worksheet"]] if source.get("worksheet") in wb.sheetnames else wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if kind == "gsheet":
        import gspread
        gc = gspread.service_account(filename=source["credentials"],
                                     scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
        ws = gc.open_by_key(source["sheet_id"]).worksheet(source["worksheet"])
        return ws.get_all_values()
    raise ValueError(f"Unknown source type {kind!r}")
